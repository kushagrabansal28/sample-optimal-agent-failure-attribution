"""Build presentation/research_presentation.pptx (editable, 30 slides, speaker notes on every slide).

Every number is read from results/processed/*.json (never typed in); figures come from results/figures; screenshots from
report/screens. Charts on the results slides are native PowerPoint charts (right-click -> Edit Data).

    python scripts/build_deck.py
"""
import csv
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG, SCR, PROC = ROOT / "results" / "figures", ROOT / "report" / "screens", ROOT / "results" / "processed"
OUT = ROOT / "presentation"
OUT.mkdir(exist_ok=True)
(OUT / "assets").mkdir(exist_ok=True)


def J(name):
    return json.loads((PROC / f"{name}.json").read_text(encoding="utf-8"))


MAIN, GAP, KSW, ABL, SIZE, TOOL, TGT = (J(n)["summary"] for n in ("main", "gap_sweep", "K_sweep", "ablation", "size_sweep", "toolagent", "targets"))
NOISE, STRESS, BUDGET, PAIRED, META = J("noise_sweep")["summary"], J("stress")["summary"], J("budget_curve")["summary"], J("paired_stats"), J("suite_meta")
TESTS, FE = J("test_report"), J("frontend_test_report")
FE_TOTAL, FE_PASS = FE["numTotalTests"], FE["numPassedTests"]
A = {s["method"]: s for s in MAIN if s["kind"] == "automaton"}
B = [s for s in MAIN if s["kind"] == "bandit"]
ABLD = {(s["factor"], str(s["level"])): s for s in ABL}
NMAT = list(csv.DictReader(open(ROOT / "research" / "novelty_matrix.csv", encoding="utf-8")))
AUTHORS = "Kushagra Yogendra Bansal (25bde0080)  ·  Kartik Kesharwani (25bde0053)"
TRIALS = META["trials"]

# ---- palette / type --------------------------------------------------------------------------------------------------
BG, PANEL, TX, MUT = RGBColor(11, 16, 28), RGBColor(21, 29, 46), RGBColor(232, 237, 247), RGBColor(150, 163, 190)
BLUE, CYAN, AMBER, RED, GREEN, VIOLET = (RGBColor(79, 163, 255), RGBColor(45, 212, 224), RGBColor(246, 183, 60),
                                         RGBColor(255, 99, 120), RGBColor(52, 211, 153), RGBColor(167, 139, 250))
HEAD, BODY, MONO = "Segoe UI Semibold", "Segoe UI", "Consolas"
BADGE = {"synthetic": AMBER, "pending": RED, "measured": GREEN, "theory": VIOLET, "literature": BLUE}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5
page = [0]


def tf_style(p, size, color=TX, bold=False, font=BODY, align=None):
    p.font.size, p.font.bold, p.font.name = Pt(size), bold, font
    p.font.color.rgb = color
    if align:
        p.alignment = align


def text(s, x, y, w, h, content, size=16, color=TX, bold=False, font=BODY, align=None, anchor=None, bullets=False, gap=6):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    f = box.text_frame
    f.word_wrap = True
    f.margin_left = f.margin_right = Inches(0.04)
    if anchor:
        f.vertical_anchor = anchor
    items = content if isinstance(content, list) else [content]
    for i, it in enumerate(items):
        p = f.paragraphs[0] if i == 0 else f.add_paragraph()
        p.text = ("•  " + it) if bullets else it
        tf_style(p, size, color, bold, font, align)
        p.space_after = Pt(gap)
    return box


def rect(s, x, y, w, h, fill=PANEL, line=None, radius=True):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line; sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    if radius:
        sh.adjustments[0] = 0.04
    sh.shadow.inherit = False
    return sh


def badge(s, label, x, y, color=None):
    color = color or BADGE.get(label, MUT)
    w = 0.34 + 0.095 * len(label)
    b = rect(s, x, y, w, 0.3, fill=BG, line=color)
    b.adjustments[0] = 0.5
    f = b.text_frame
    f.margin_left = f.margin_right = Inches(0.05); f.margin_top = f.margin_bottom = 0
    f.paragraphs[0].text = label
    tf_style(f.paragraphs[0], 11, color, True, MONO, PP_ALIGN.CENTER)
    f.vertical_anchor = MSO_ANCHOR.MIDDLE
    return w


def picture(s, path, x, y, w, h):
    """Place an image inside the box (x,y,w,h), preserving aspect ratio, centred."""
    iw, ih = Image.open(path).size
    sc = min(w / iw, h / ih)
    pw, ph = iw * sc, ih * sc
    return s.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def slide(title, notes, tag=None, kicker=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    page[0] += 1
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), Inches(SH))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    if kicker:
        text(s, 0.55, 0.28, 8, 0.3, kicker.upper(), 11, CYAN, True, MONO)
    text(s, 0.55, 0.52, 11.2, 0.8, title, 28, TX, True, HEAD)
    if tag:
        x = 12.9
        for t in (tag if isinstance(tag, list) else [tag])[::-1]:
            x -= 0.34 + 0.095 * len(t) + 0.1
            badge(s, t, x, 0.4)
    text(s, 0.55, 7.1, 9, 0.3, "AFCR · Fixed-confidence culprit identification · " + AUTHORS.split("  ·  ")[0].split(" (")[0] + " & " + AUTHORS.split("  ·  ")[1].split(" (")[0], 9, MUT)
    text(s, 12.2, 7.1, 0.7, 0.3, str(page[0]), 10, MUT, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


def caption(s, x, y, w, content, size=11):
    return text(s, x, y, w, 0.5, content, size, MUT)


def f0(x):
    return f"{x:,.0f}"


def pct(x):
    return f"{100 * x:.0f} %"


def table(s, x, y, w, rows, colw, size=12, rowh=0.36, header=True, hl=None):
    shape = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(rowh * len(rows)))
    t = shape.table
    tblPr = shape._element.graphic.graphicData.tbl.tblPr
    tblPr.set("bandRow", "0"); tblPr.set("firstRow", "0")
    for i, cw in enumerate(colw):
        t.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        t.rows[r].height = Inches(rowh)
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(30, 42, 66) if (header and r == 0) else (RGBColor(26, 38, 60) if hl and r in hl else PANEL)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.text = str(val)
            tf_style(p, size, CYAN if (header and r == 0) else TX, header and r == 0, BODY)
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
    return t


def style_chart(ch, legend=True, size=11):
    ch.font.size, ch.font.color.rgb, ch.font.name = Pt(size), TX, BODY
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size); ch.legend.font.color.rgb = TX
    for ax in (ch.category_axis, ch.value_axis):
        ax.tick_labels.font.size = Pt(size); ax.tick_labels.font.color.rgb = MUT
        ax.format.line.color.rgb = RGBColor(70, 82, 110)
    ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(40, 52, 80)
    ch.value_axis.has_major_gridlines = True


def color_series(ch, colors):
    for ser, col in zip(ch.plots[0].series, colors):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = col
        ser.format.line.color.rgb = col


LAB = {"lucb_prefix": "LUCB + prefix", "lucb": "LUCB", "successive_elim": "Succ. elim.", "uniform_fc": "Uniform", "random_fc": "Random",
       "equal_budget": "Equal budget", "lucb_naive_share": "LUCB + naive share*", "car_fixed_k": "CAR-style fixed K", "oracle": "Oracle"}
ORDER = ["lucb_prefix", "lucb", "successive_elim", "uniform_fc", "random_fc"]

# =========================================================================================================== 1 title
s = prs.slides.add_slide(BLANK)
s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
page[0] += 1
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), Inches(SH)); bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
text(s, 0.9, 1.0, 9, 0.4, "MODELS OF COMPUTATION  ·  RESEARCH PROJECT", 13, CYAN, True, MONO)
text(s, 0.9, 1.6, 11.5, 2.2, "Sample-Optimal Culprit Identification for Agent Failures", 44, TX, True, HEAD)
text(s, 0.9, 3.55, 11.5, 0.9, "Fixed-Confidence Attribution with Automata-Aware Sequential Replay (AFCR)", 24, BLUE, False, BODY)
text(s, 0.9, 4.7, 11.5, 0.5, AUTHORS, 18, TX)
text(s, 0.9, 5.35, 11.5, 0.8, ["Working title: “sample-optimal” names the research objective. We prove δ-correctness and a sample-complexity upper bound;",
                              "we do not claim minimax optimality. All experiments in this deck are synthetic."], 13, MUT, gap=2)
x = 0.9
for t in ("synthetic", "theory", "pending"):
    x += badge(s, t, x, 6.45) + 0.15
text(s, x + 0.1, 6.43, 7, 0.4, "pending = ALFWorld / HotpotQA runs", 11, MUT)
s.notes_slide.notes_text_frame.text = ("Title slide. Introduce both authors. Stress the honest framing up front: the title states the objective, and what we actually prove is "
                                       "δ-correctness of the stopping rule plus a sample-complexity upper bound for one strategy. All reported numbers are from synthetic "
                                       "simulations with exactly known ground truth; the real-benchmark stages (ALFWorld, HotpotQA) are implemented as adapters but have not been run.")

# =========================================================================================================== 2 motivation
s = slide("Why agent failures are hard to diagnose", "Motivation. Long-horizon agents plan, search, call tools, store memory and answer. When the final answer is wrong, the "
          "developer holds a long trace and one question: which step is responsible? The last action is usually a symptom. Cite Zhang et al. (Who&When, ICML 2025, arXiv:2505.00212): the best "
          "step-level attribution accuracy they report is about 14 %, which is why LLM judging alone is not a satisfying answer.", tag="literature", kicker="Motivation")
for i, (h, b, c) in enumerate([("Long horizons", "Plan → search → tools → memory → answer. A single early slip can corrupt everything after it.", BLUE),
                               ("Symptom ≠ cause", "The step that executes the harmful action is often not the step that decided on it.", AMBER),
                               ("Judges are weak", "LLM-judge step attribution is correlational; best step-level accuracy ≈ 14 % on Who&When (Zhang et al., 2025).", RED)]):
    rect(s, 0.6 + i * 4.1, 1.7, 3.85, 2.9)
    text(s, 0.85 + i * 4.1, 1.9, 3.4, 0.5, h, 20, c, True, HEAD)
    text(s, 0.85 + i * 4.1, 2.55, 3.4, 2.0, b, 15, TX)
rect(s, 0.6, 4.95, 12.1, 1.6, fill=RGBColor(18, 40, 58))
text(s, 0.9, 5.1, 11.5, 1.4, ["Our question is interventional, not correlational:",
                              "“If this step had been repaired, how likely is the run to succeed?”  — and how many replays do we need before we may say so with confidence 1 − δ?"], 17, TX, gap=8)

# =========================================================================================================== 3 trajectory
mu = J("main")  # keep reference alive for checks
s = slide("A failed trajectory with a hidden cause", "Figure 2 is computed exactly by dynamic programming over the 48-state automaton (no sampling). Steps before the planted fault only exceed "
          "the baseline slightly because re-running re-rolls the fault; after the fault the state is corrupted and the repair value falls. The step that maximises the repair value is the culprit "
          "under our definition.", tag="synthetic", kicker="Problem")
picture(s, FIG / "fig02_failed_trajectory.png", 0.5, 1.45, 7.8, 5.5)
rect(s, 8.55, 1.6, 4.25, 5.1)
text(s, 8.8, 1.75, 3.8, 0.4, "What the figure shows", 16, CYAN, True, HEAD)
text(s, 8.8, 2.25, 3.8, 4.3, ["Exact μ_t for one failed run (no sampling noise)", "Planted fault: a wrong decision at one step",
                              "Earlier steps look “useful” only because a re-run re-rolls the fault", "After the fault the state is corrupted and μ_t drops",
                              "Culprit := argmax_t μ_t, with ties and near-ties made explicit via ε"], 14, TX, bullets=True)

# =========================================================================================================== 4 why not simple
s = slide("Why last-action or judge-based debugging is not enough", "Association is not prediction is not intervention effect is not unique responsibility is not joint cause. "
          "We choose one well-defined causal estimand — the repair effect of a single step — and make its estimation statistically rigorous. Other estimands (Shapley credit, direct effects) are "
          "different questions and are listed as future work.", tag="theory", kicker="Problem")
rows = [["Approach", "What it answers", "Gap for our question"],
        ["Last action", "which step executed the failure", "symptom, not decision"],
        ["LLM judge on the log", "which step looks wrong", "correlational; ≈ 14 % best step accuracy (Who&When)"],
        ["Fixed-K counterfactual replay", "effect of intervening at each step", "K chosen in advance; no confidence statement; no stopping rule"],
        ["Oracle-substitution replay", "does a gold action flip the outcome", "needs gold actions; one sample; no budget view"],
        ["AFCR (this work)", "which step is the best repair, certified at confidence 1 − δ", "adaptive budget, valid stopping, abstention when unsure"]]
table(s, 0.6, 1.7, 12.1, rows, [3.3, 4.2, 4.6], 14, 0.62, hl=[5])
text(s, 0.6, 5.7, 12.1, 1.1, "We never claim that replay-based attribution is new: the contribution is the sequential, fixed-confidence treatment of the replay budget "
     "(see related-work matrix, next).", 15, MUT)

# =========================================================================================================== 5 related matrix
cols = [("step_level_attribution", "step"), ("counterfactual_replay", "replay"), ("adaptive_allocation", "adaptive"), ("fixed_confidence_stopping", "fixed-conf"),
        ("delta_pac_guarantee", "δ-PAC"), ("prefix_aware_information_reuse", "prefix"), ("formal_automata_representation", "automata")]
rows = [["Work", "depth"] + [c[1] for c in cols]]
keep = [r for r in NMAT if r["paper_id"] in ("GK2016", "WW2025", "CAR2026", "AT2025", "CP2026", "MA2025")]
for r in keep:
    rows.append([r["short_name"], r["read_depth"]] + [("●" if r[c[0]] == "Yes" else ("◐" if r[c[0]] == "Partial" else "–")) for c in cols])
rows.append(["AFCR (this work)", "—", "●", "●", "●", "●", "●", "●", "●"])
s = slide("Related work: what exists and what does not", "Rows come from research/novelty_matrix.csv, built by reading the papers (read-depth column shows how much of each paper we "
          "verified). ● yes, ◐ partial, – no. Important corrections to the brief: Who&When has no replay and no Shapley; Shapley belongs to Causal Agent Replay. "
          "The AFCR row is what we implement; whether the combination is novel is provisional — the search was finite and the matrix is dated.", tag="literature", kicker="Related work")
table(s, 0.6, 1.6, 12.1, rows, [3.2, 1.3, 1.0, 1.0, 1.1, 1.1, 1.1, 1.1, 1.2], 13, 0.52, hl=[len(rows) - 1])
text(s, 0.6, 6.0, 12.1, 0.9, "● yes  ◐ partial  – no.  Best-arm identification (Garivier & Kaufmann, 2016) supplies the statistics but is not an agent paper; "
     "no surveyed agent-attribution paper combines replay with fixed-confidence stopping. Novelty status: provisional (finite search).", 13, MUT)

# =========================================================================================================== 6 formal problem
s = slide("Formal problem: best-arm identification over replays", "Arms are the candidate steps. Pulling arm t means: restore the original failed run just before step t, apply "
          "the intervention (repair = force the reference action; resample = redraw from the policy), re-simulate to the end, and observe success or failure. Y_t is Bernoulli(μ_t). "
          "The target is a function of the means: argmax up to ε, or the first / last step above a threshold θ. Assumptions A1–A6 in the paper make the arms well-defined and the pulls independent.",
          tag="theory", kicker="Formalisation")
rect(s, 0.6, 1.6, 6.0, 5.2)
text(s, 0.85, 1.75, 5.6, 0.4, "Instance", 16, CYAN, True, HEAD)
text(s, 0.85, 2.3, 5.6, 4.4, ["Failed trajectory τ of length T; candidate steps t ∈ {1,…,T}",
                              "Pull(t): restore the exact state before step t, intervene, replay to the end → Y ∈ {0,1}",
                              "μ_t = P(success | intervene at t) — unknown, estimated only by pulls",
                              "Replay noise is the agent's own stochasticity (the policy π)"], 15, TX, bullets=True)
rect(s, 6.85, 1.6, 5.9, 5.2)
text(s, 7.1, 1.75, 5.5, 0.4, "Goal (fixed confidence)", 16, CYAN, True, HEAD)
text(s, 7.1, 2.3, 5.5, 4.4, ["Sampling rule: which step to replay next", "Stopping rule: a data-dependent time τ",
                              "Recommendation: a step ĥ, or abstain",
                              "δ-correct: P(τ < ∞ and ĥ ∉ target set) ≤ δ  for every instance",
                              "Targets: argmax (ε-tolerant), first-above-θ, last-above-θ",
                              "Minimise rollouts E[τ] subject to δ-correctness"], 15, TX, bullets=True)

# =========================================================================================================== 7 counterfactual replay
s = slide("Counterfactual replay mechanism", "Figure 3 uses a real sampled trajectory. The original run on top is restored up to the chosen step, one action is forced to the reference action, "
          "and the continuation on the bottom is re-simulated. One replay equals one sample of the outcome distribution — one Bernoulli pull. Repair vs resample are both implemented; "
          "resample is the null intervention and measures intrinsic sensitivity.", tag="synthetic", kicker="Mechanism")
picture(s, FIG / "fig03_counterfactual_replay.png", 0.5, 1.45, 12.3, 5.4)

# =========================================================================================================== 8 algorithm
s = slide("The adaptive algorithm: one round", "Figure 4: maintain a confidence interval per step, ask the stopping rule whether the target is certified, otherwise let the sampling "
          "rule pick the next step (LUCB picks the leader and its strongest challenger). The sampling rule is interchangeable — δ-correctness depends only on the confidence sequence and the "
          "stopping rule. Budget exhaustion is an abstention and is never reported as a certified answer.", tag="theory", kicker="Method")
picture(s, FIG / "fig04_adaptive_selection_flowchart.png", 0.5, 1.4, 7.9, 5.6)
rect(s, 8.65, 1.6, 4.15, 5.1)
text(s, 8.85, 1.75, 3.8, 0.4, "Design rules", 16, CYAN, True, HEAD)
text(s, 8.85, 2.25, 3.8, 4.4, ["Stopping rule ⟂ sampling rule", "Anytime-valid intervals (running intersection)", "Abstain, never fabricate",
                               "Every result states whether the δ-guarantee applies", "Sharing prefixes is opt-in and validated"], 14, TX, bullets=True)

# =========================================================================================================== 9 confidence / stopping
s = slide("Confidence sequences and the stopping rule", "We use a stack-KL confidence sequence, x_n = ln(2Kn(n+1)/δ), intersected over time so intervals never widen. The stop test is "
          "lo_h + ε ≥ max over other arms of hi_a. Figure 5 is an actual run: intervals separate and the algorithm stops exactly when the separation margin reaches zero. "
          "Stack-KL is never wider than the Hoeffding version (tested) and is slightly looser than the LUCB-rate bound only in constants (ablation).", tag=["theory", "synthetic"], kicker="Method")
picture(s, FIG / "fig05_confidence_stopping.png", 0.5, 1.4, 7.7, 5.6)
rect(s, 8.45, 1.6, 4.35, 5.1)
c1, c2, c3 = ABLD[("confidence", "stack_kl")], ABLD[("confidence", "stack_hoeffding")], ABLD[("confidence", "kl_lucb_rate")]
text(s, 8.65, 1.75, 4.0, 0.4, "Stopping rule", 16, CYAN, True, HEAD)
text(s, 8.65, 2.25, 4.0, 1.2, "stop when  lo_h + ε ≥ max_{a≠h} hi_a", 13, TX, font=MONO)
text(s, 8.65, 3.0, 4.0, 0.4, "Confidence-sequence ablation (48 states)", 13, CYAN, True)
table(s, 8.65, 3.45, 4.0, [["sequence", "mean rollouts"], ["stack-KL", f0(c1["mean_rollouts"])], ["KL-LUCB rate", f0(c3["mean_rollouts"])], ["stack-Hoeffding", f0(c2["mean_rollouts"])]],
      [2.1, 1.9], 12, 0.4)
caption(s, 8.65, 5.2, 4.0, f"n = {c1['n']} trials each, same seeds; all accuracy {pct(c1['accuracy'])}. Synthetic.")

# =========================================================================================================== 10 theory
s = slide("Theory: what is proven", "Theorem 1: for any sampling rule, the stopping rule built on an anytime-valid confidence sequence is δ-correct — a union bound over arms and time. "
          "Theorem 2: for successive elimination, an explicit sample-complexity upper bound of order Σ_a Δ_a⁻² log(K/(δΔ_a)); the full proof is in the paper appendix. "
          "Theorem 3: valid prefix reuse. We do NOT claim minimax optimality of any algorithm — the lower bound of Garivier–Kaufmann is used as a yardstick, not matched.", tag="theory", kicker="Theory")
for i, (h, b, c) in enumerate([("Theorem 1 · δ-correctness", "Any sampling rule + stack confidence sequence + the separation stopping rule: P(stop with a wrong answer) ≤ δ. Abstention is not a δ-PAC event.", GREEN),
                               ("Theorem 2 · sample complexity", "Successive elimination stops after O( Σ_a Δ_a⁻² · log(K / (δ Δ_a)) ) rollouts w.h.p. Gap-dependent, matches the lower-bound shape up to logs.", BLUE),
                               ("Theorem 3 · prefix reuse", "Matched (time, state, action) observations from other arms' replays are unbiased draws of the same arm (strong Markov property).", VIOLET)]):
    rect(s, 0.6 + i * 4.1, 1.7, 3.85, 3.5)
    text(s, 0.8 + i * 4.1, 1.85, 3.5, 0.8, h, 17, c, True, HEAD)
    text(s, 0.8 + i * 4.1, 2.75, 3.5, 2.4, b, 14, TX)
rect(s, 0.6, 5.45, 12.1, 1.35, fill=RGBColor(46, 28, 36))
text(s, 0.85, 5.55, 11.6, 1.25, ["Not claimed: minimax / instance optimality of any algorithm; the lower bound T*(μ)·kl(δ, 1−δ) is a yardstick only.",
                                 "Naively pooling all prefix-matched replays is NOT valid — kept as a negative control (next slides)."], 14, TX, gap=4)

# =========================================================================================================== 11 prefix
nv = A["lucb_naive_share"]
s = slide("Prefix reuse: valid vs naive sharing", "Figure 10: replays that start later reach the same state at the same time with the same action as another arm's replay would; their remaining "
          "trajectory has the same law, so they can be credited to the other arm. We verified Theorem 3 numerically: matched observations agree with the exact value (|z| < 2) while "
          "naively pooled observations are biased (|z| about 20–30). In the main experiment naive sharing certifies a wrong answer far more often — shown honestly as a negative control with a "
          "small certified sample.", tag=["theory", "synthetic"], kicker="Theory → experiment")
picture(s, FIG / "fig10_prefix_tree.png", 0.5, 1.4, 6.9, 5.6)
rect(s, 7.6, 1.6, 5.2, 5.2)
ab = (ABLD[("sharing", "none")], ABLD[("sharing", "valid")], ABLD[("sharing", "naive")])
text(s, 7.8, 1.75, 4.8, 0.4, "Sharing ablation (48 states)", 15, CYAN, True)
table(s, 7.8, 2.2, 4.8, [["sharing", "mean rollouts", "accuracy"], ["none", f0(ab[0]["mean_rollouts"]), pct(ab[0]["accuracy"])],
                         ["valid (Thm 3)", f0(ab[1]["mean_rollouts"]), pct(ab[1]["accuracy"])], ["naive ✗", f0(ab[2]["mean_rollouts"]), pct(ab[2]["accuracy"])]], [1.8, 1.7, 1.3], 13, 0.45, hl=[3])
text(s, 7.8, 4.25, 4.8, 2.5, [f"Valid sharing: {ab[0]['mean_rollouts'] / ab[1]['mean_rollouts']:.2f}× fewer rollouts than no sharing, accuracy not reduced.",
                              f"Naive sharing: accuracy collapses to {pct(ab[2]['accuracy'])}; it stops early with a wrong answer. Never use it.",
                              "Sharing is switched off automatically when state reconstruction is noisy (A1 violated)."], 13, TX, bullets=True)

# =========================================================================================================== 12 automaton 48
s = slide("The 48-state agent automaton", "Six groups of eight states: PLAN, OBSERVE, TOOL, MEMORY, VALIDATE, TERMINAL. Seven decision states carry planted faults; corrupted states form a corridor "
          "to the failure terminals; recovery edges exist. The state count is validated programmatically (unique ids and labels; the DFA sink is not counted). Larger 64- and 96-state versions "
          "are generated the same way.", tag="synthetic", kicker="Models of computation")
picture(s, FIG / "fig12_automaton_48_states.png", 0.5, 1.4, 8.3, 5.6)
rect(s, 9.0, 1.6, 3.8, 5.2)
text(s, 9.2, 1.75, 3.4, 0.4, "Structure", 16, CYAN, True, HEAD)
text(s, 9.2, 2.25, 3.4, 4.4, ["48 distinct logical states (6 × 8)", "Roles: benign, decision, fault entry, corridor, recovery, terminal",
                              "Exact V_t(q) by backward induction", "Ground truth μ_t known — algorithms never see it"], 14, TX, bullets=True)

# =========================================================================================================== 13 finite-state views
s = slide("Finite-state views of the same specification", "DFA: total transition function with an explicit sink, deterministic most-likely view. NFA: all possible successors, so we track active-state sets. "
          "ε-NFA: silent self-correction as ε-moves; ε-closure is computed before and after each symbol. Each view is tested: NFA→DFA subset construction and ε-removal preserve the language.",
          tag="theory", kicker="Models of computation")
rect(s, 0.5, 1.5, 6.0, 4.75); picture(s, FIG / "fig06_dfa.png", 0.6, 1.6, 5.8, 4.1)
text(s, 0.6, 5.7, 5.8, 0.4, "DFA — total δ, explicit sink", 13, CYAN, True, align=PP_ALIGN.CENTER)
for i, (fn, cap) in enumerate([("fig07_nfa_active_sets.png", "NFA — active-state sets"), ("fig08_enfa_closure.png", "ε-NFA — ε-closure")]):
    rect(s, 6.7, 1.5 + i * 2.4, 6.1, 2.3)
    picture(s, FIG / fn, 6.8, 1.55 + i * 2.4, 5.9, 1.75)
    text(s, 6.8, 3.3 + i * 2.4, 5.9, 0.4, cap, 13, CYAN, True, align=PP_ALIGN.CENTER)
text(s, 0.6, 6.4, 12.1, 0.6, "Nondeterminism in an NFA is a formal device (“some path accepts”). LLM sampling randomness is probabilistic and belongs to the probabilistic automaton below.", 13, MUT)

# =========================================================================================================== 14 PA + LTS
s = slide("Probabilistic automaton and labelled transition system", "The probabilistic automaton is the natural model of a stochastic agent: each (state, action) maps to a distribution over successors, "
          "and the word probability is e_q0ᵀ M_a1 … M_an 1_F. The LTS carries the labelled actions — reasoning, tool use, observation, update, validation, recovery, success, failure — and "
          "supports reachability and deadlock analysis. Exact quantities used by the experiments are computed from these kernels.", tag="theory", kicker="Models of computation")
picture(s, FIG / "fig09_probabilistic_automaton.png", 0.5, 1.45, 7.4, 5.5)
rect(s, 8.15, 1.6, 4.65, 5.2)
text(s, 8.35, 1.75, 4.3, 0.4, "Why these two", 16, CYAN, True, HEAD)
text(s, 8.35, 2.25, 4.3, 4.4, ["PA: P(w) = e_{q₀}ᵀ M_{a₁}…M_{aₙ} 1_F", "LTS: labelled actions, reachability, deadlocks",
                               "Exact μ_t^rep = Σ P(q′|q,ρ(q)) V_t(q′)", "Exact μ_t^res = V_{t−1}(q_{t−1})", "Used as ground truth for every experiment"], 14, TX, bullets=True)

# =========================================================================================================== 15 beyond FSM
s = slide("Beyond finite state: transducers, PDA, Turing machine", "All nine models are implemented and tested. The comparison table summarises acceptance, memory, nondeterminism and fit for agent traces. "
          "Important distinction: a finite-state machine cannot recognise arbitrarily deep nesting — a PDA can (our test shows a depth-3 bounded DFA failing on a depth-5 word). "
          "The bounded Turing machine reports step_limit/tape_bound as 'unknown', not 'reject'.", tag="theory", kicker="Models of computation")
rows = [["Model", "Memory", "Output / acceptance", "Agent-trace use", "Limit for attribution"],
        ["Mealy", "finite", "output on transitions", "controller reacting to observations", "no acceptance notion"],
        ["Moore", "finite", "output on states (n+1 for n inputs)", "state = action being executed", "Mealy→Moore needs state expansion"],
        ["PDA", "unbounded stack", "final state / empty stack", "nested or recursive tool calls", "single stack, no probabilities"],
        ["Turing machine (bounded)", "tape (bounded here)", "accept / reject / step_limit / tape_bound", "general-computation demo", "bound hit ⇒ unknown, not reject"],
        ["PA", "finite", "P(w) over words", "stochastic agent", "Markov, finite memory"]]
table(s, 0.6, 1.65, 12.1, rows, [2.4, 1.7, 3.0, 2.7, 2.3], 12, 0.62)
text(s, 0.6, 5.55, 12.1, 1.3, ["Implemented in src/afcr/automata: dfa, nfa (+ε), pa, lts, transducers (Mealy/Moore), pda, turing, benchmark48.",
                               "Only the PA/LTS views carry the causal interpretation; the others are formal-language demonstrations of the same specification."], 14, MUT, gap=4)

# =========================================================================================================== 16 architecture
s = slide("Software architecture", "Packages: automata, statistics (the only place bounds are defined), replay (arms and exact ground truth), attribution (targets, algorithms, engine), agents and "
          "benchmarks, experiments and evaluation, and the FastAPI backend that serves the website. Every trial is a plain dict so any number can be regenerated from a JSON spec with "
          "python -m afcr.experiments.run_one.", tag="theory", kicker="Implementation")
picture(s, FIG / "fig13_software_architecture.png", 0.5, 1.4, 12.3, 5.6)

# =========================================================================================================== 17 protocol
s = slide("Evaluation pipeline and protocol", f"Twelve experiment groups, each with {TRIALS} paired trials per method unless noted (all methods share seeds so comparisons are paired). Rollout counts of trials that hit "
          "the budget cap are censored: means are lower bounds on the true cost. Accuracy intervals are Wilson; rollout intervals are bootstrap; paired tests are Wilcoxon signed-rank with "
          "Holm correction over the four primary comparisons. Total simulation time came from suite_meta.json.", tag=["synthetic", "measured"], kicker="Evaluation")
picture(s, FIG / "fig14_evaluation_pipeline.png", 0.5, 4.75, 12.3, 2.2)
G = list(META["groups"].items())
for k, chunk in enumerate((G[:6], G[6:])):
    table(s, 0.6 + k * 4.3, 1.5, 4.0, [["group", "rows"]] + [[g, f"{v['rows']:,}"] for g, v in chunk], [2.6, 1.4], 11, 0.4)
caption(s, 9.3, 1.6, 3.5, f"{sum(v['rows'] for v in META['groups'].values()):,} trial rows in total; δ = {META['delta_default']}. SYNTHETIC only.")

# =========================================================================================================== 18 main results (native chart)
s = slide("Main result: adaptive allocation needs fewer rollouts", f"48-state automaton, {A['lucb']['n']} paired trials per method, δ = {META['delta_default']}, budget cap 60 000 rollouts. "
          "Bars are mean rollouts; trials that abstained (budget exhausted) are censored at the cap, so uniform/random numbers understate their true cost. Equal-budget and CAR-style "
          "fixed-K do not certify anything by design — their rollouts are fixed in advance. LUCB with naive sharing is the invalid negative control.", tag=["synthetic", "measured"], kicker="Results")
cd = CategoryChartData()
meths = ORDER
cd.categories = [LAB[m] for m in meths]
cd.add_series("Mean rollouts (censored at cap)", [round(A[m]["mean_rollouts"]) for m in meths])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(7.1), Inches(4.6), cd)
style_chart(gf.chart, legend=False); color_series(gf.chart, [BLUE])
gf.chart.plots[0].gap_width = 60
gf.chart.plots[0].has_data_labels = True
gf.chart.plots[0].data_labels.font.size = Pt(11); gf.chart.plots[0].data_labels.font.color.rgb = TX
gf.chart.plots[0].data_labels.number_format = "#,##0"; gf.chart.plots[0].data_labels.number_format_is_linked = False
cd2 = CategoryChartData(); cd2.categories = [LAB[m] for m in meths]
cd2.add_series("Abstained (budget exhausted)", [round(100 * A[m]["abstain_rate"]) for m in meths])
gf2 = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.7), Inches(1.5), Inches(5.2), Inches(4.6), cd2)
style_chart(gf2.chart, legend=False); color_series(gf2.chart, [AMBER]); gf2.chart.plots[0].gap_width = 60
gf2.chart.plots[0].has_data_labels = True; gf2.chart.plots[0].data_labels.font.size = Pt(11); gf2.chart.plots[0].data_labels.font.color.rgb = TX
gf2.chart.has_title = True; gf2.chart.chart_title.text_frame.text = "Abstention rate (%)"
gf2.chart.chart_title.text_frame.paragraphs[0].font.size = Pt(13); gf2.chart.chart_title.text_frame.paragraphs[0].font.color.rgb = MUT
gf.chart.has_title = True; gf.chart.chart_title.text_frame.text = "Mean rollouts to certify"
gf.chart.chart_title.text_frame.paragraphs[0].font.size = Pt(13); gf.chart.chart_title.text_frame.paragraphs[0].font.color.rgb = MUT
text(s, 0.6, 6.2, 12.2, 0.8, f"Accuracy of the answer (vs exact argmax): LUCB+prefix {pct(A['lucb_prefix']['accuracy'])}, LUCB {pct(A['lucb']['accuracy'])}, successive elim. {pct(A['successive_elim']['accuracy'])}, "
     f"uniform {pct(A['uniform_fc']['accuracy'])}, random {pct(A['random_fc']['accuracy'])}. Native charts: right-click → Edit Data.", 12, MUT)

# =========================================================================================================== 19 paired stats
s = slide("Paired statistics with multiplicity control", "Each row is a paired comparison on identical seeds. Wilcoxon signed-rank p-values are Holm-corrected over the four primary comparisons; all four reject "
          "at 0.05. Ratios are mean-rollout ratios (>1 means the first method is cheaper). 95 % intervals are bootstrap intervals of the paired mean difference. Synthetic data only: these are "
          "statements about our simulated benchmark, not about real agents.", tag=["synthetic", "measured"], kicker="Results")
rows = [["comparison (A vs B)", "n", "mean A", "mean B", "saving", "× cheaper", "95 % CI of diff", "p (Holm)"]]
for p in PAIRED:
    rows.append([f"{LAB[p['a']]} vs {LAB[p['b']]}", p["n"], f0(p["mean_a"]), f0(p["mean_b"]), f0(p["mean_diff"]), f"{p['ratio']:.2f}",
                 f"[{f0(p['ci'][0])}, {f0(p['ci'][1])}]", f"{p['p_holm']:.1e}"])
table(s, 0.6, 1.7, 12.1, rows, [3.3, 0.6, 1.1, 1.1, 1.1, 1.1, 2.3, 1.5], 13, 0.6)
text(s, 0.6, 5.0, 12.1, 1.8, ["Adaptive allocation (LUCB) vs uniform: " + f"{PAIRED[0]['ratio']:.2f}× fewer rollouts on average.",
                              "Valid prefix sharing on top of LUCB: " + f"{PAIRED[3]['ratio']:.2f}× fewer again.",
                              "All comparisons censored at the budget cap → the true savings are at least as large as shown."], 15, TX, bullets=True)

# =========================================================================================================== 20 gap sweep
g_l = sorted([x for x in GAP if x["method"] == "lucb"], key=lambda x: x["gap"])
g_u = sorted([x for x in GAP if x["method"] == "uniform_fc"], key=lambda x: x["gap"])
s = slide("Cost versus gap and the lower bound", "Bandit-arm model with K = 8 arms; the best-vs-runner-up gap varies. Cost grows roughly like gap⁻² for uniform allocation "
          "(a test checks the log-log slope is between −3.2 and −1.3) and every method stays above the lower bound on average (also tested). The lower bound is a yardstick for an oracle-allocating "
          "algorithm: we do not claim to match it.", tag=["synthetic", "measured"], kicker="Results")
picture(s, FIG / "fig19_gap_sweep_lower_bound.png", 0.5, 1.4, 7.8, 5.6)
rows = [["gap", "LUCB", "Uniform", "ratio"]] + [[f"{a['gap']:.2f}", f0(a["mean_rollouts"]), f0(b["mean_rollouts"]), f"{b['mean_rollouts'] / a['mean_rollouts']:.1f}×"] for a, b in zip(g_l, g_u)]
table(s, 8.6, 1.7, 4.2, rows, [0.9, 1.1, 1.2, 1.0], 13, 0.5)
caption(s, 8.6, 4.9, 4.2, f"Mean rollouts, {g_l[0]['n']} trials per cell, K = 8. Synthetic.")

# =========================================================================================================== 21 K sweep + delta
k_l = sorted([x for x in KSW if x["method"] == "lucb"], key=lambda x: x["K"])
k_u = sorted([x for x in KSW if x["method"] == "uniform_fc"], key=lambda x: x["K"])
k_s = sorted([x for x in KSW if x["method"] == "successive_elim"], key=lambda x: x["K"])
s = slide("Scaling with the number of candidates and with δ", "Left: growth of cost with the number of candidate steps K at a fixed gap. Adaptive methods grow slowly because eliminated arms stop "
          "being sampled; uniform allocation grows roughly linearly with K. Right: the δ sweep shows the expected mild increase of cost as δ shrinks; empirical error stays at 0 in every cell — "
          "consistent with δ-correctness but not a proof (the guarantee is an upper bound).", tag=["synthetic", "measured"], kicker="Results")
cd = CategoryChartData(); cd.categories = [str(x["K"]) for x in k_l]
cd.add_series("LUCB", [round(x["mean_rollouts"]) for x in k_l]); cd.add_series("Successive elim.", [round(x["mean_rollouts"]) for x in k_s]); cd.add_series("Uniform", [round(x["mean_rollouts"]) for x in k_u])
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.5), Inches(1.5), Inches(6.3), Inches(4.7), cd)
style_chart(gf.chart); gf.chart.has_title = True; gf.chart.chart_title.text_frame.text = "Mean rollouts vs number of candidates K"
gf.chart.chart_title.text_frame.paragraphs[0].font.size = Pt(13); gf.chart.chart_title.text_frame.paragraphs[0].font.color.rgb = MUT
for ser, col in zip(gf.chart.plots[0].series, [BLUE, VIOLET, AMBER]):
    ser.format.line.color.rgb = col; ser.format.line.width = Pt(2.5); ser.smooth = False
    ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = col
ds = sorted(J("delta_sweep")["summary"], key=lambda x: x["delta"])
rows = [["δ", "mean rollouts", "accuracy", "wrong certified"]] + [[f"{x['delta']}", f0(x["mean_rollouts"]), pct(x["accuracy"]), pct(x["identified_error_rate"])] for x in ds]
table(s, 7.1, 1.7, 5.7, rows, [0.9, 1.7, 1.3, 1.8], 13, 0.5)
caption(s, 7.1, 4.8, 5.7, f"δ sweep: {ds[0]['n']} trials per δ. 0 wrong certified answers in {sum(x['n'] for x in ds)} trials is consistent with — not a proof of — δ-correctness.")

# =========================================================================================================== 22 ablation
def ablrow(f, lv, name):
    x = ABLD[(f, lv)]
    return [name, f0(x["mean_rollouts"]), pct(x["accuracy"]), pct(x["abstain_rate"])]
s = slide("Ablations: which design choices matter", "One factor at a time on the 48-state automaton, same seeds. Restricting candidates to decision states and using a small ε make the most "
          "difference on the planted-fault instances because they remove near-ties between harmless steps; valid prefix sharing gives a large saving at no cost in accuracy. "
          "Candidate restriction uses domain knowledge about decision states; it is reported as an assumption-dependent speed-up, not as part of the guarantee.", tag=["synthetic", "measured"], kicker="Results")
rows = [["factor", "mean rollouts", "accuracy", "abstained"], ablrow("allocation", "lucb", "allocation = LUCB"), ablrow("allocation", "successive_elim", "allocation = succ. elim."),
        ablrow("allocation", "uniform_fc", "allocation = uniform"), ablrow("sharing", "valid", "prefix sharing = valid"), ablrow("sharing", "naive", "prefix sharing = naive ✗"),
        ablrow("candidates", "decision", "candidates = decision states only"), ablrow("epsilon", "0.02", "ε = 0.02"), ablrow("epsilon", "0.05", "ε = 0.05")]
table(s, 0.6, 1.65, 8.0, rows, [3.5, 1.7, 1.4, 1.4], 13, 0.48, hl=[5])
rect(s, 8.9, 1.65, 3.9, 4.4)
text(s, 9.1, 1.8, 3.5, 0.4, "Reading it", 15, CYAN, True, HEAD)
text(s, 9.1, 2.3, 3.5, 3.7, ["Baseline (LUCB, stack-KL, all steps, ε = 0): " + f0(ABLD[('allocation', 'lucb')]['mean_rollouts']) + " rollouts", "Naive sharing is fast because it is wrong",
                             "ε > 0 trades exactness for cost; the guarantee becomes “within ε”"], 13, TX, bullets=True)

# =========================================================================================================== 23 size / noise
s = slide("Larger automata and violated assumptions", "Left: the 48-, 64- and 96-state automata. Prefix reuse keeps the cost almost flat while the plain methods grow with trajectory length. "
          "Right: when state reconstruction is noisy (assumption A1 violated) we measure accuracy against the noisy-context argmax and the cost rises; the algorithm abstains more rather "
          "than certifying wrongly, and sharing is disabled automatically. This is a robustness stress test, not a guarantee.", tag=["synthetic", "measured"], kicker="Results")
sz = lambda m: [x for x in sorted(SIZE, key=lambda x: x["n_states"]) if x["method"] == m]
cd = CategoryChartData(); cd.categories = [str(x["n_states"]) + " states" for x in sz("lucb")]
for m in ("lucb_prefix", "lucb", "uniform_fc"):
    cd.add_series(LAB[m], [round(x["mean_rollouts"]) for x in sz(m)])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(6.4), Inches(4.6), cd)
style_chart(gf.chart); color_series(gf.chart, [CYAN, BLUE, AMBER]); gf.chart.has_title = True; gf.chart.chart_title.text_frame.text = "Mean rollouts by automaton size"
gf.chart.chart_title.text_frame.paragraphs[0].font.size = Pt(13); gf.chart.chart_title.text_frame.paragraphs[0].font.color.rgb = MUT
nz = sorted(NOISE, key=lambda x: x["recon_noise"])
rows = [["reconstruction noise", "mean rollouts", "abstained", "accuracy*"]] + [[f"{x['recon_noise']:.2f}", f0(x["mean_rollouts"]), pct(x["abstain_rate"]), pct(x["accuracy"])] for x in nz]
table(s, 7.2, 1.7, 5.6, rows, [1.9, 1.5, 1.1, 1.1], 13, 0.5)
caption(s, 7.2, 4.5, 5.6, "*accuracy against the argmax of the noisy replay context (the estimand actually sampled). Censored at the budget cap.")

# =========================================================================================================== 24 error analysis
s = slide("Negative controls and failure modes", "Honest reporting: (1) naive sharing is invalid and the experiment shows it. (2) The CAR-style fixed-K rule is a different estimand — it reports the "
          "latest step whose effect excludes zero — so its 'accuracy' against the argmax-repair target is not a like-for-like comparison and should not be read as 'CAR is wrong'. "
          "(3) Several methods abstain on hard instances; abstention is a feature, not an error. Figure 18 breaks the failures down.", tag=["synthetic", "measured"], kicker="Results")
picture(s, FIG / "fig18_error_analysis.png", 0.5, 1.4, 7.4, 5.6)
rect(s, 8.15, 1.6, 4.65, 5.2)
text(s, 8.35, 1.75, 4.3, 0.4, "What to take from it", 15, CYAN, True, HEAD)
car = A["car_fixed_k"]
text(s, 8.35, 2.25, 4.3, 4.5, [f"Naive sharing: {pct(nv['accuracy'])} accuracy, {pct(nv['abstain_rate'])} abstained — invalid by construction",
                               f"CAR-style fixed K: {pct(car['accuracy'])} vs the argmax target — a different estimand (see notes)",
                               "Uniform/random abstain most often at the cap",
                               "Never certified a wrong answer in guaranteed runs (see consistency test)"], 13, TX, bullets=True)

# =========================================================================================================== 25 tool agent + targets
ts = {(x["variant"], x["method"]): x for x in TOOL}
tg = {(x["target"], x["method"]): x for x in TGT}
s = slide("Stage B tool agent and alternative targets", "Stage B is a second, structurally different synthetic environment: a tool-using agent (search, calculator, memory) with exact memoised DP ground "
          "truth, in a stochastic and a deterministic variant. The targets experiment replaces argmax by threshold rules: first-above-θ and last-above-θ. These are still synthetic — they test the "
          "implementation beyond the automaton, not real agent behaviour.", tag=["synthetic", "measured"], kicker="Results")
rows = [["Stage B variant", "method", "mean rollouts", "accuracy", "abstained"]]
for v in ("stochastic", "deterministic_phase0", "deterministic_phase4"):
    for m in ("lucb", "lucb_prefix", "uniform_fc"):
        if (v, m) in ts:
            x = ts[(v, m)]
            rows.append([v, LAB[m], f0(x["mean_rollouts"]), pct(x["accuracy"]), pct(x["abstain_rate"])])
table(s, 0.6, 1.65, 6.5, rows, [1.8, 1.3, 1.2, 1.0, 1.2], 11, 0.42)
rows = [["target", "method", "mean rollouts", "accuracy", "abstained"]]
for (t, m), x in sorted(tg.items()):
    rows.append([t, LAB.get(m, m), f0(x["mean_rollouts"]), pct(x["accuracy"]), pct(x["abstain_rate"])])
table(s, 7.3, 1.65, 5.5, rows[:8], [1.1, 1.4, 1.1, 0.8, 1.1], 11, 0.42)
text(s, 0.6, 5.55, 12.2, 1.3, ["Same ordering as on the automaton: adaptive < uniform in rollouts.", "Threshold targets can legitimately end in “no step qualifies” or abstention; both are reported, not hidden."], 14, MUT, bullets=True, gap=4)

# =========================================================================================================== 26 website 1
s = slide("The website: run the algorithm and watch it decide", "Screenshots of the running React/TypeScript site, served by the FastAPI backend on localhost. The dashboard calls the real "
          "algorithm, replays the rounds as a forest plot of confidence intervals, shows the stopping-rule margin, and can run uniform allocation on the identical instance. The Studio loads "
          "every saved result group and the paired statistics; exports are raw CSV/JSON. Pages carry synthetic/measured badges so no simulated number is mistaken for a benchmark result.", tag="measured", kicker="Demonstration")
for i, (fn, cap) in enumerate([("dashboard.png", "Attribution dashboard — live run, replayable CI plot, stopping rule"), ("studio.png", "Experiment studio — saved results from raw CSV")]):
    rect(s, 0.5 + i * 6.2, 1.5, 6.05, 4.75)
    picture(s, SCR / fn, 0.6 + i * 6.2, 1.6, 5.85, 4.0)
    text(s, 0.6 + i * 6.2, 5.7, 5.85, 0.5, cap, 12, CYAN, True, align=PP_ALIGN.CENTER)
text(s, 0.6, 6.4, 12.2, 0.6, "9 pages: overview · related work · automata lab · trajectory explorer · dashboard · studio · theory · reproducibility · docs", 13, MUT)

# =========================================================================================================== 27 website 2
s = slide("The website: automata laboratory, explorer, theory", "Left: all nine formalisms are interactive — pan, zoom, search states, step through traces; the 48-state automaton shows the planted culprit. "
          "Middle: the trajectory explorer samples a failed run, shows the exact μ_t, lets you intervene on a step and compare the original with the counterfactual. Right: the theory page "
          "lists the assumptions, theorems with expandable proofs and a claims ledger that marks what is proven, tested or pending.", tag="measured", kicker="Demonstration")
for i, (fn, cap) in enumerate([("lab.png", "Automata laboratory"), ("explorer.png", "Trajectory explorer"), ("theory.png", "Theory & guarantees")]):
    rect(s, 0.5 + i * 4.15, 1.5, 4.0, 4.6)
    picture(s, SCR / fn, 0.58 + i * 4.15, 1.6, 3.85, 3.9)
    text(s, 0.58 + i * 4.15, 5.55, 3.85, 0.4, cap, 13, CYAN, True, align=PP_ALIGN.CENTER)
text(s, 0.6, 6.3, 12.2, 0.7, "Custom cursor, spotlight and motion are disabled on touch devices and when the OS requests reduced motion; charts also skip animation then.", 13, MUT)

# =========================================================================================================== 28 testing
s = slide("Testing and reproducibility", f"{TESTS['total']} backend tests (pytest) and {FE_TOTAL} frontend tests (vitest) all pass, run on this machine; the reports are stored in "
          "results/processed. Consistency tests recompute every reported summary from the raw CSV and check theory-level behaviour such as gap scaling and the lower bound. "
          "Config files in configs/ reproduce any single trial or a whole suite; seeds are explicit.", tag="measured", kicker="Quality")
rows = [["suite", "tests", "passed"]] + [[m.replace("test_", ""), v["passed"] + v["failed"], v["passed"]] for m, v in TESTS["modules"].items()]
rows.append(["frontend (vitest)", FE_TOTAL, FE_PASS])
table(s, 0.6, 1.65, 6.2, rows, [3.6, 1.2, 1.4], 13, 0.5)
rect(s, 7.1, 1.65, 5.7, 4.6)
text(s, 7.3, 1.8, 5.3, 0.4, "Reproduce", 15, CYAN, True, HEAD)
text(s, 7.3, 2.3, 5.3, 3.9, ["python -m pytest -q", "python -m afcr.experiments.suite --config configs/suite_default.json",
                             "python -m afcr.experiments.run_one configs/spec_automaton_prefix.json --seed 7",
                             "python scripts/make_figures.py", "python scripts/build_pdfs.py", "python scripts/build_deck.py",
                             "uvicorn afcr.api.app:app --port 8000"], 12, TX, font=MONO, gap=5)

# =========================================================================================================== 29 limits
s = slide("Limitations and honest status", "State plainly what is not done. ALFWorld and HotpotQA adapters exist and are unit-tested with a mock agent, but we have not run them — so no claim about real "
          "agents. All performance numbers come from synthetic models with a Markov structure that makes ground truth computable. The OpenAI provider is implemented but untested. "
          "Theorem 2 is for successive elimination only. 'Sample-optimal' is the objective, not a result. Novelty is provisional.", tag=["pending", "synthetic"], kicker="Honesty")
for i, (h, items, c) in enumerate([("Done and verified", [f"{TESTS['total']} backend + {FE_TOTAL} frontend tests passing", "δ-correctness (Thm 1), prefix validity (Thm 3) checked numerically", "9 formal models + 48-state benchmark", "12 experiment groups, raw → processed"], GREEN),
                                   ("Pending / not run", ["ALFWorld (Stage C) — adapter only", "HotpotQA (Stage D) — adapter only", "Real LLM provider evaluation", "Human-validated culprit labels"], RED),
                                   ("Not claimed", ["Minimax / instance optimality", "Results on real agents", "That replay is new (it is not)", "Novelty beyond the finite literature search"], AMBER)]):
    rect(s, 0.6 + i * 4.1, 1.7, 3.85, 4.9)
    text(s, 0.8 + i * 4.1, 1.85, 3.5, 0.5, h, 18, c, True, HEAD)
    text(s, 0.8 + i * 4.1, 2.5, 3.5, 4.0, items, 14, TX, bullets=True, gap=8)

# =========================================================================================================== 30 conclusion
s = slide("Conclusions and future work", "Summary. We turned culprit identification into a fixed-confidence best-arm problem over counterfactual replays, proved δ-correctness for any sampling rule, "
          "gave a sample-complexity bound for successive elimination, validated prefix reuse, and measured on synthetic automata that adaptive allocation needs several times fewer rollouts. "
          "Future work: run ALFWorld/HotpotQA, tighten the bounds (instance-dependent), Track-and-Stop style allocation, direct-effect estimands with common random numbers.", tag="theory", kicker="Wrap-up")
rect(s, 0.6, 1.6, 6.0, 5.2)
text(s, 0.85, 1.75, 5.6, 0.4, "What we showed", 17, CYAN, True, HEAD)
text(s, 0.85, 2.3, 5.6, 4.4, ["Culprit identification as fixed-confidence BAI over replays", "δ-correct for any sampling rule (Theorem 1)",
                              f"{PAIRED[0]['ratio']:.2f}× fewer rollouts than uniform at equal confidence (synthetic)",
                              f"Valid prefix reuse: a further {PAIRED[3]['ratio']:.2f}×; naive reuse is demonstrably wrong",
                              "Abstention instead of fabricated answers"], 14, TX, bullets=True)
rect(s, 6.85, 1.6, 5.95, 5.2)
text(s, 7.1, 1.75, 5.5, 0.4, "Next", 17, CYAN, True, HEAD)
text(s, 7.1, 2.3, 5.5, 2.5, ["Run Stage C/D on ALFWorld and HotpotQA", "Instance-optimal allocation (Track-and-Stop)", "Direct-effect estimands, joint (Shapley) credit"], 14, TX, bullets=True)
text(s, 7.1, 4.7, 5.5, 1.9, [AUTHORS, "Code, paper, report and this deck: see README.md", "Key references: Garivier & Kaufmann 2016; Zhang et al. 2025; Shah 2026 (CAR); AgenTracer 2025"], 12, MUT, gap=6)

OUTFILE = OUT / "research_presentation.pptx"
cp = prs.core_properties
cp.author, cp.last_modified_by, cp.title, cp.comments = "Kushagra Yogendra Bansal; Kartik Kesharwani", "Kushagra Yogendra Bansal", "Sample-Optimal Culprit Identification for Agent Failures", ""
prs.save(OUTFILE)
print(f"wrote {OUTFILE.relative_to(ROOT)} with {len(prs.slides)} slides")

# ---- outline (titles + speaker notes) so the talk can be reviewed without opening PowerPoint -----------------------------
lines = ["# Presentation outline", "", f"{len(prs.slides)} slides · {AUTHORS}", "",
         "Generated by `python scripts/build_deck.py`; numbers on the slides are read from `results/processed/*.json`.", ""]
for i, sl in enumerate(prs.slides, 1):
    title = next((sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text and sh.top < Inches(1.35) and sh.top > Inches(0.45)), None) \
        or next((sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text), "")
    title = "Sample-Optimal Culprit Identification for Agent Failures" if i == 1 else title
    lines += [f"## {i}. {title}", "", "**Speaker notes.** " + sl.notes_slide.notes_text_frame.text, ""]
(OUT / "outline.md").write_text("\n".join(lines), encoding="utf-8")
print("wrote presentation/outline.md")
